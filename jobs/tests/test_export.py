"""Phan logic cua Export Job — chay duoc ma khong can BigQuery.

Ba thu de sai nhat deu nam o day: file mang dung so cua nguon, dau ban ky
noi that ve mon no, va cat sheet khi vuot gioi han cua Excel.
"""

from datetime import datetime, timezone

from conftest import export_main

EXCEL_MAX_ROWS = export_main.EXCEL_MAX_ROWS
stamp_lines = export_main.stamp_lines
to_rows = export_main.to_rows
write_csv = export_main.write_csv


class Row:
    """Gia lap mot dong tra ve tu BigQuery."""

    def __init__(self, year, state, gender, name, number, market_share):
        self.year, self.state, self.gender = year, state, gender
        self.name, self.number, self.market_share = name, number, market_share


def ban_ky(**kw):
    base = {
        "id": 7, "signed_version_id": 3, "label": "Bao cao Q3",
        "signed_by": "lead@cty.com",
        "signed_at": datetime(2026, 9, 21, 9, 30, tzinfo=timezone.utc),
        "source_run_ids": ["run-A", "run-B"], "checksum": "abc123",
        "rules_version": 2, "violations": None, "violations_fingerprint": None,
        "open_tickets": None, "approval_note": None,
    }
    base.update(kw)
    return base


# ------------------------------------------------- so trong file la so nguon

def test_file_mang_dung_so_cua_nguon():
    """Khong con duong nao sua so giua nguon va file.

    Day la ly do ca lop override bi bo: truoc kia file ban ra va BigQuery
    co the khac nhau ma khong ai phat hien.
    """
    out = list(to_rows([Row(2021, "CA", "F", "Emma", 60, 0.6),
                        Row(2021, "CA", "F", "Olivia", 40, 0.4)]))
    assert [r[4] for r in out] == [60, 40]
    assert [r[5] for r in out] == ["0.6000000000", "0.4000000000"]
    # Khong con cot `da_sua`: khong co gi de danh dau nua.
    assert len(out[0]) == len(export_main.HEADER) == 6


def test_thi_phan_rong_khi_nguon_khong_co():
    out = list(to_rows([Row(2021, "CA", "F", "Emma", 60, None)]))
    assert out[0][5] == ""


def test_ghi_csv_du_dong_va_du_tieu_de(tmp_path):
    rows = list(to_rows([Row(2021, "CA", "F", "Emma", 60, 0.6)]))
    path = tmp_path / "x.csv"
    n, canh_bao = write_csv(str(path), rows)

    assert n == 1 and canh_bao is None
    dong = path.read_text().splitlines()
    assert dong[0] == "year,state,gender,name,number,market_share"
    assert dong[1] == "2021,CA,F,Emma,60,0.6000000000"


def test_csv_khong_bi_nhet_dau_ban_ky_vao_giua_du_lieu(tmp_path):
    """Mot dong chu thich o dau file lam Excel doc lech cot.

    Dau phai nam o file rieng, nen write_csv nhan `stamp` roi bo qua.
    """
    rows = list(to_rows([Row(2021, "CA", "F", "Emma", 60, 0.6)]))
    path = tmp_path / "x.csv"
    write_csv(str(path), rows, stamp_lines(ban_ky(), None, 1))

    dong = path.read_text().splitlines()
    assert len(dong) == 2
    assert dong[0].startswith("year,")


# -------------------------------------------------------------- dau ban ky

def test_dau_ban_ky_noi_that_khi_ban_sach():
    txt = "\n".join(stamp_lines(ban_ky(), ["CA"], 120))
    assert "Bao cao Q3" in txt
    assert "run-A, run-B" in txt
    assert "abc123" in txt
    assert "khong con vi pham luat nao" in txt.lower()
    assert "khong con ticket nao chua dong" in txt.lower()


def test_dau_ban_ky_khong_giau_mon_no():
    """Ban ky kem no va ban sach khong duoc trong giong nhau."""
    job = ban_ky(violations={"duoi_nguong_kiem_duyet": 3, "bien_dong_bat_thuong": 405},
                 violations_fingerprint="f00d", open_tickets=[123, 140],
                 approval_note="3 o duoi nguong la so that cua bang nho")
    txt = "\n".join(stamp_lines(job, ["CA", "TX"], 9))

    assert "VAN CON vi pham" in txt
    assert "bien_dong_bat_thuong: 405" in txt
    assert "duoi_nguong_kiem_duyet: 3" in txt
    assert "#123, #140" in txt
    assert "3 o duoi nguong la so that cua bang nho" in txt
    assert "f00d" in txt


def test_dau_ban_ky_ghi_ro_pham_vi():
    assert "CA, TX" in "\n".join(stamp_lines(ban_ky(), ["CA", "TX"], 1))
    assert "tat ca cac bang" in "\n".join(stamp_lines(ban_ky(), None, 1))


# ------------------------------------------------------------------- xlsx

def test_xlsx_co_sheet_bia_ban_ky_dung_truoc_du_lieu(tmp_path):
    """Excel mo dung sheet dau — nguoi mo file thay minh dang cam ban nao
    truoc khi thay so."""
    from openpyxl import load_workbook

    job = ban_ky(violations={"tang_dot_bien": 23}, open_tickets=[5])
    rows = list(to_rows([Row(2021, "CA", "F", "Emma", 60, 0.6)]))
    path = tmp_path / "x.xlsx"
    n, _ = export_main.write_xlsx(str(path), rows, stamp_lines(job, None, 1))

    assert n == 1
    wb = load_workbook(path)
    assert wb.sheetnames[0] == "Ban ky"
    assert wb.sheetnames[1] == "Trang 1"
    bia = "\n".join(str(c[0].value or "") for c in wb["Ban ky"].iter_rows())
    assert "tang_dot_bien: 23" in bia
    assert "#5" in bia


def test_xlsx_vuot_gioi_han_thi_tach_sheet_chu_khong_cat_bot(tmp_path, monkeypatch):
    """Gia gioi han xuong 5 dong de khoi phai sinh mot trieu dong that."""
    monkeypatch.setattr(export_main, "EXCEL_MAX_ROWS", 5)
    from openpyxl import load_workbook

    rows = [[2021, "CA", "F", f"Ten{i}", i, "0.1"] for i in range(12)]
    path = tmp_path / "x.xlsx"
    n, canh_bao = export_main.write_xlsx(str(path), rows)

    assert n == 12, "khong duoc cat bot dong nao"
    assert canh_bao and "tach thanh 3 sheet" in canh_bao

    wb = load_workbook(path)
    assert wb.sheetnames == ["Trang 1", "Trang 2", "Trang 3"]
    # Moi sheet co tieu de rieng, 4 dong du lieu, 4 dong, 4 dong
    assert [ws.max_row for ws in wb.worksheets] == [5, 5, 5]
    assert wb["Trang 2"].cell(row=1, column=1).value == "year"


def test_gioi_han_that_la_cua_excel():
    assert EXCEL_MAX_ROWS == 1_048_576
